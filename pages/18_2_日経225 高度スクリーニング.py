import datetime
import io
import os
import time
import warnings
import numpy as np
import pandas as pd
import streamlit as st
import yfinance as yf

warnings.filterwarnings("ignore")

# ============================================================
# 保存先フォルダ・ファイルの設定 (日本株用)
# ============================================================
DATA_DIR = "data_japan"
DATA_FILE_RAW = os.path.join(DATA_DIR, "nikkei225_raw_data.csv")
DATA_FILE_JP = os.path.join(DATA_DIR, "nikkei225_japanese_spreadsheet.csv")

# ============================================================
# 初期設定 (Streamlit画面の設定)
# ============================================================
st.set_page_config(page_title="日経225 高度スクリーニング", layout="wide")

DEFAULT_CONFIG = {
    "nikkei225_url": "https://en.wikipedia.org/wiki/Nikkei_225",
    "top_n": 50,
    "request_sleep_seconds": 0.15,
    "missing_data_passes": False,
    "excluded_sectors": [],
    # 日本企業の実態に最適化したデフォルト財務フィルター基準
    "filters": {
        "revenue_jpy": {
            "min": 100_000_000_000,
            "max": None,
        },  # 売上高1,000億円以上
        "revenue_growth_1y": {"min": 0.05, "max": None},  # 売上成長率 5%以上
        "revenue_cagr_3y": {"min": 0.03, "max": None},  # 3年CAGR 3%以上
        "net_margin": {"min": 0.08, "max": None},  # 純利益率 8%以上
        "roe": {
            "min": 0.08,
            "max": None,
        },  # ROE 8%以上 (伊藤レポート目標基準)
        "fcf_margin": {"min": 0.05, "max": None},  # FCF利益率 5%以上
        "net_debt_to_ebitda": {
            "min": None,
            "max": 3.00,
        },  # ネット負債/EBITDA 3倍以下
        "trailing_pe": {"min": 0, "max": 35},  # 実績PER 35倍以下
        "forward_pe": {"min": 0, "max": 30},  # 予想PER 30倍以下
        "price_to_book": {"min": None, "max": None},
        "price_to_sales": {"min": None, "max": 10},
        "ev_to_ebitda": {"min": 0, "max": 20},
        "peg_ratio": {"min": 0, "max": None},
    },
    "quality_score_metrics": {
        "revenue_growth_1y": {"direction": "high", "weight": 1.0},
        "revenue_cagr_3y": {"direction": "high", "weight": 1.0},
        "net_margin": {"direction": "high", "weight": 1.0},
        "roe": {"direction": "high", "weight": 1.0},
        "fcf_margin": {"direction": "high", "weight": 1.0},
        "net_debt_to_ebitda": {"direction": "low", "weight": 0.5},
    },
    "valuation_score_metrics": {
        "trailing_pe": {"direction": "low", "weight": 1.0},
        "forward_pe": {"direction": "low", "weight": 1.0},
        "price_to_book": {"direction": "low", "weight": 0.5},
        "price_to_sales": {"direction": "low", "weight": 0.5},
        "ev_to_ebitda": {"direction": "low", "weight": 1.0},
        "peg_ratio": {"direction": "low", "weight": 0.5},
    },
    "score_weights": {"quality": 0.60, "valuation": 0.40},
    "sort_by": "composite_score",
    "sort_ascending": False,
}

# ============================================================
# スプレッドシート用 日本語カラム対応表 (日本株用)
# ============================================================
JAPANESE_COLUMNS_MAP = {
    "code": "銘柄コード",
    "symbol": "ティッカー",
    "company_name": "企業名",
    "sector": "業種",
    "current_price": "株価(円)",
    "price_above_sma200_label": "200日線上",
    "pct_from_52w_high": "52週高値乖離率(%)",
    "beta": "ベータ値(Beta)",
    "market_cap_jpy": "時価総額(円)",
    "enterprise_value_jpy": "企業価値EV(円)",
    "revenue_jpy": "売上高(円)",
    "net_income_jpy": "純利益(円)",
    "ebitda_jpy": "EBITDA(円)",
    "free_cash_flow_jpy": "FCF(円)",
    "revenue_growth_1y": "売上高成長率(%)",
    "revenue_cagr_3y": "売上高3年CAGR(%)",
    "operating_margin": "営業利益率(%)",
    "net_margin": "純利益率(%)",
    "fcf_margin": "FCF利益率(%)",
    "roe": "ROE(%)",
    "total_debt_jpy": "総負債(円)",
    "cash_jpy": "保有現金(円)",
    "net_debt_jpy": "ネット有利子負債(円)",
    "net_debt_to_ebitda": "ネット負債/EBITDA(倍)",
    "trailing_pe": "実績PER(倍)",
    "forward_pe": "予想PER(倍)",
    "price_to_book": "PBR(倍)",
    "price_to_sales": "PSR(倍)",
    "ev_to_ebitda": "EV/EBITDA(倍)",
    "peg_ratio": "PEGレシオ",
    "dividend_yield": "配当利回り(%)",
    "quality_score": "クオリティスコア",
    "valuation_score": "割安スコア",
    "composite_score": "総合スコア",
}


# ============================================================
# ローカルファイル保存・読み込み関数
# ============================================================
def save_data_to_local(df):
  os.makedirs(DATA_DIR, exist_ok=True)
  df.to_csv(DATA_FILE_RAW, index=False)
  jp_df = create_spreadsheet_export_dataframe(df)
  jp_df.to_csv(DATA_FILE_JP, index=False, encoding="utf-8-sig")


def load_data_from_local():
  if os.path.exists(DATA_FILE_RAW):
    try:
      return pd.read_csv(DATA_FILE_RAW)
    except Exception:
      return None
  return None


def get_saved_data_time():
  if os.path.exists(DATA_FILE_RAW):
    mtime = os.path.getmtime(DATA_FILE_RAW)
    return datetime.datetime.fromtimestamp(mtime).strftime("%Y/%m/%d %H:%M:%S")
  return None


# ============================================================
# データ取得系
# ============================================================
@st.cache_data(ttl=3600 * 24)
def get_nikkei225_constituents(url):
  """日経225の構成銘柄一覧を取得し、4桁コード＋.T に整形"""
  storage_options = {
      "User-Agent": (
          "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML,"
          " like Gecko) Chrome/120.0.0.0 Safari/537.36"
      )
  }
  tables = pd.read_html(url, storage_options=storage_options)

  target_table = None
  for t in tables:
    cols = [str(c).strip().lower() for c in t.columns]
    has_code = any(
        k in cols
        for k in ["symbol", "ticker", "code", "コード", "銘柄コード"]
    )
    has_company = any(
        k in cols for k in ["company", "security", "社名", "銘柄名", "企業名"]
    )
    if has_code and has_company:
      target_table = t.copy()
      break

  if target_table is None:
    for t in tables:
      if len(t) >= 200:
        target_table = t.copy()
        break

  if target_table is None:
    target_table = tables[0].copy()

  code_col, company_col, sector_col = None, None, None
  for col in target_table.columns:
    c_lower = str(col).strip().lower()
    if not code_col and any(
        k in c_lower
        for k in ["symbol", "ticker", "code", "コード", "銘柄コード"]
    ):
      code_col = col
    elif not company_col and any(
        k in c_lower for k in ["company", "security", "社名", "銘柄名", "企業名"]
    ):
      company_col = col
    elif not sector_col and any(
        k in c_lower for k in ["sector", "業種", "industry"]
    ):
      sector_col = col

  if not code_col:
    code_col = target_table.columns[0]
  if not company_col:
    company_col = (
        target_table.columns[1] if len(target_table.columns) > 1 else code_col
    )
  if not sector_col:
    sector_col = (
        target_table.columns[2]
        if len(target_table.columns) > 2
        else company_col
    )

  df = pd.DataFrame()
  df["code"] = (
      target_table[code_col].astype(str).str.extract(r"(\d{4})")[0].fillna("")
  )
  df = df[df["code"] != ""].copy()
  df["Security"] = target_table.loc[df.index, company_col].astype(str)
  df["Sector"] = target_table.loc[df.index, sector_col].astype(str)

  # 日本株はコード末尾に .T を付与（例: 7203.T）
  df["yahoo_symbol"] = df["code"] + ".T"
  df["display_symbol"] = df["code"] + ".T"

  return df.drop_duplicates(subset=["code"]).reset_index(drop=True)


def get_statement_series(statement, possible_names):
  if statement is None or statement.empty:
    return pd.Series(dtype="float64")
  for name in possible_names:
    if name in statement.index:
      values = pd.to_numeric(statement.loc[name], errors="coerce").dropna()
      return values.sort_index(ascending=False)
  return pd.Series(dtype="float64")


def first_valid_value(series):
  if series is None or len(series) == 0:
    return np.nan
  value = series.iloc[0]
  return np.nan if pd.isna(value) else float(value)


def safe_float(value):
  try:
    if value is None or pd.isna(value):
      return np.nan
    return float(value)
  except (TypeError, ValueError):
    return np.nan


def safe_divide(numerator, denominator):
  if pd.isna(numerator) or pd.isna(denominator) or denominator == 0:
    return np.nan
  return numerator / denominator


def get_first_info_value(info, names):
  for name in names:
    value = safe_float(info.get(name))
    if not pd.isna(value):
      return value
  return np.nan


def get_market_cap(ticker, info):
  try:
    value = safe_float(ticker.fast_info["market_cap"])
    if not pd.isna(value):
      return value
  except Exception:
    pass
  return get_first_info_value(info, ["marketCap"])


def analyze_company(company):
  yahoo_symbol = company["yahoo_symbol"]
  ticker = yf.Ticker(yahoo_symbol)

  income_statement = ticker.get_income_stmt(freq="yearly")
  balance_sheet = ticker.get_balance_sheet(freq="yearly")
  cash_flow = ticker.get_cash_flow(freq="yearly")

  try:
    info = ticker.get_info()
  except Exception:
    info = {}

  revenue = get_statement_series(
      income_statement, ["TotalRevenue", "OperatingRevenue"]
  )
  net_income = get_statement_series(income_statement, [
      "NetIncomeCommonStockholders",
      "NetIncome",
      "NetIncomeIncludingNoncontrollingInterests",
  ])
  ebitda = get_statement_series(income_statement, ["EBITDA", "NormalizedEBITDA"])
  operating_income = get_statement_series(
      income_statement, ["OperatingIncome"]
  )

  equity = get_statement_series(balance_sheet, [
      "StockholdersEquity",
      "CommonStockEquity",
      "TotalEquityGrossMinorityInterest",
  ])
  total_debt = get_statement_series(balance_sheet, ["TotalDebt"])
  cash = get_statement_series(balance_sheet, [
      "CashCashEquivalentsAndShortTermInvestments",
      "CashAndCashEquivalents",
      "CashFinancial",
  ])

  free_cash_flow = get_statement_series(cash_flow, ["FreeCashFlow"])
  operating_cash_flow = get_statement_series(
      cash_flow, ["OperatingCashFlow", "TotalCashFromOperatingActivities"]
  )
  capital_expenditure = get_statement_series(
      cash_flow, ["CapitalExpenditure", "CapitalExpenditures"]
  )

  latest_revenue = first_valid_value(revenue)
  latest_net_income = first_valid_value(net_income)
  latest_ebitda = first_valid_value(ebitda)
  latest_operating_income = first_valid_value(operating_income)
  latest_equity = first_valid_value(equity)
  latest_total_debt = first_valid_value(total_debt)
  latest_cash = first_valid_value(cash)
  latest_fcf = first_valid_value(free_cash_flow)
  latest_ocf = first_valid_value(operating_cash_flow)
  latest_capex = first_valid_value(capital_expenditure)

  if (
      pd.isna(latest_fcf)
      and not pd.isna(latest_ocf)
      and not pd.isna(latest_capex)
  ):
    latest_fcf = (
        latest_ocf + latest_capex
        if latest_capex < 0
        else latest_ocf - latest_capex
    )

  revenue_growth_1y = np.nan
  if len(revenue) >= 2:
    prev_rev = safe_float(revenue.iloc[1])
    if prev_rev > 0:
      revenue_growth_1y = latest_revenue / prev_rev - 1

  revenue_cagr_3y = np.nan
  if len(revenue) >= 4:
    rev_3y = safe_float(revenue.iloc[3])
    if latest_revenue > 0 and rev_3y > 0:
      revenue_cagr_3y = (latest_revenue / rev_3y) ** (1 / 3) - 1

  net_margin = safe_divide(latest_net_income, latest_revenue)
  operating_margin = safe_divide(latest_operating_income, latest_revenue)
  fcf_margin = safe_divide(latest_fcf, latest_revenue)

  roe = np.nan
  if len(equity) >= 2:
    prev_eq = safe_float(equity.iloc[1])
    if latest_equity > 0 and prev_eq > 0:
      roe = safe_divide(latest_net_income, (latest_equity + prev_eq) / 2)

  net_debt = np.nan
  net_debt_to_ebitda = np.nan
  if not pd.isna(latest_total_debt):
    cash_val = 0 if pd.isna(latest_cash) else latest_cash
    net_debt = latest_total_debt - cash_val
    if not pd.isna(latest_ebitda) and latest_ebitda > 0:
      net_debt_to_ebitda = net_debt / latest_ebitda

  market_cap = get_market_cap(ticker, info)

  # 日本株テクニカル指標
  current_price = safe_float(
      info.get("currentPrice") or info.get("regularMarketPrice")
  )
  sma200 = safe_float(info.get("twoHundredDayAverage"))
  high_52w = safe_float(info.get("fiftyTwoWeekHigh"))
  beta = safe_float(info.get("beta"))

  is_above_sma200 = (
      (current_price >= sma200)
      if (pd.notna(current_price) and pd.notna(sma200))
      else False
  )
  pct_from_52w_high = (
      ((current_price / high_52w) - 1)
      if (pd.notna(current_price) and pd.notna(high_52w) and high_52w > 0)
      else np.nan
  )

  return {
      "code": company["code"],
      "symbol": company["display_symbol"],
      "company_name": company["Security"],
      "sector": company["Sector"],
      "current_price": current_price,
      "sma200": sma200,
      "is_above_sma200": is_above_sma200,
      "price_above_sma200_label": "○" if is_above_sma200 else "×",
      "pct_from_52w_high": pct_from_52w_high,
      "beta": beta,
      "market_cap_jpy": market_cap,
      "enterprise_value_jpy": get_first_info_value(info, ["enterpriseValue"]),
      "revenue_jpy": latest_revenue,
      "net_income_jpy": latest_net_income,
      "ebitda_jpy": latest_ebitda,
      "free_cash_flow_jpy": latest_fcf,
      "revenue_growth_1y": revenue_growth_1y,
      "revenue_cagr_3y": revenue_cagr_3y,
      "operating_margin": operating_margin,
      "net_margin": net_margin,
      "fcf_margin": fcf_margin,
      "roe": roe,
      "total_debt_jpy": latest_total_debt,
      "cash_jpy": latest_cash,
      "net_debt_jpy": net_debt,
      "net_debt_to_ebitda": net_debt_to_ebitda,
      "trailing_pe": get_first_info_value(info, ["trailingPE"]),
      "forward_pe": get_first_info_value(info, ["forwardPE"]),
      "price_to_book": get_first_info_value(info, ["priceToBook"]),
      "price_to_sales": get_first_info_value(
          info, ["priceToSalesTrailing12Months"]
      ),
      "ev_to_ebitda": get_first_info_value(info, ["enterpriseToEbitda"]),
      "peg_ratio": get_first_info_value(info, ["trailingPegRatio", "pegRatio"]),
      "dividend_yield": get_first_info_value(info, ["dividendYield"]),
  }


# ============================================================
# スコア計算・フィルタリング
# ============================================================
def clean_valuation_values(dataframe):
  val_cols = [
      "trailing_pe",
      "forward_pe",
      "price_to_book",
      "price_to_sales",
      "ev_to_ebitda",
      "peg_ratio",
  ]
  for col in val_cols:
    if col in dataframe.columns:
      dataframe.loc[dataframe[col] <= 0, col] = np.nan
  return dataframe


def calculate_metric_score(series, direction):
  numeric = pd.to_numeric(series, errors="coerce")
  if direction == "high":
    return numeric.rank(pct=True, ascending=True) * 100
  elif direction == "low":
    return numeric.rank(pct=True, ascending=False) * 100
  return np.nan


def calculate_group_score(dataframe, metrics):
  weighted_scores, weights = [], []
  for metric, settings in metrics.items():
    if metric not in dataframe.columns:
      continue
    weight = float(settings.get("weight", 1.0))
    m_score = calculate_metric_score(
        dataframe[metric], settings.get("direction", "high")
    )
    weighted_scores.append(m_score * weight)
    weights.append(
        pd.Series(
            np.where(m_score.notna(), weight, 0), index=dataframe.index
        )
    )

  if not weighted_scores:
    return pd.Series(np.nan, index=dataframe.index)
  return sum(weighted_scores).divide(sum(weights).replace(0, np.nan))


def add_scores(dataframe, config):
  df = clean_valuation_values(dataframe.copy())
  df["quality_score"] = calculate_group_score(
      df, config["quality_score_metrics"]
  )
  df["valuation_score"] = calculate_group_score(
      df, config["valuation_score_metrics"]
  )

  q_weight = config["score_weights"]["quality"]
  v_weight = config["score_weights"]["valuation"]

  valid_q = np.where(df["quality_score"].notna(), q_weight, 0)
  valid_v = np.where(df["valuation_score"].notna(), v_weight, 0)

  df["composite_score"] = (
      (df["quality_score"] * q_weight).fillna(0)
      + (df["valuation_score"] * v_weight).fillna(0)
  ) / np.where((valid_q + valid_v) == 0, np.nan, (valid_q + valid_v))
  return df


def apply_financial_filters(dataframe, config):
  filtered = dataframe.copy()
  for metric, limits in config["filters"].items():
    if metric not in filtered.columns:
      continue
    min_v, max_v = limits.get("min"), limits.get("max")
    values = pd.to_numeric(filtered[metric], errors="coerce")
    cond = pd.Series(True, index=filtered.index)

    if min_v is not None:
      cond &= values >= min_v
    if max_v is not None:
      cond &= values <= max_v

    if min_v is not None or max_v is not None:
      cond = (
          cond | values.isna()
          if config.get("missing_data_passes")
          else cond & values.notna()
      )
    filtered = filtered[cond]
  return filtered


def format_jpy(val):
  """円単位の数値を兆円・億円・円に見やすく変換"""
  if pd.isna(val):
    return "-"
  if abs(val) >= 1_000_000_000_000:
    return f"¥{val / 1_000_000_000_000:.2f}兆"
  elif abs(val) >= 100_000_000:
    return f"¥{val / 100_000_000:,.0f}億円"
  else:
    return f"¥{val:,.0f}"


def create_display_dataframe(dataframe):
  display_cols = [
      "code",
      "company_name",
      "sector",
      "current_price",
      "price_above_sma200_label",
      "pct_from_52w_high",
      "beta",
      "market_cap_jpy",
      "revenue_growth_1y",
      "net_margin",
      "roe",
      "trailing_pe",
      "price_to_book",
      "dividend_yield",
      "quality_score",
      "valuation_score",
      "composite_score",
  ]
  display = dataframe[
      [c for c in display_cols if c in dataframe.columns]
  ].copy()

  rename_map = {
      "code": "コード",
      "company_name": "会社名",
      "sector": "業種",
      "current_price": "株価",
      "price_above_sma200_label": "200日線上",
      "pct_from_52w_high": "52週高値比",
      "beta": "ベータ値",
      "market_cap_jpy": "時価総額",
      "revenue_growth_1y": "売上成長率",
      "net_margin": "純利益率",
      "roe": "ROE",
      "trailing_pe": "実績PER",
      "price_to_book": "PBR",
      "dividend_yield": "配当利回り",
      "quality_score": "品質スコア",
      "valuation_score": "割安スコア",
      "composite_score": "総合スコア",
  }
  display = display.rename(columns=rename_map)

  if "株価" in display.columns:
    display["株価"] = display["株価"].map(
        lambda x: f"¥{x:,.0f}" if pd.notna(x) else "-"
    )
  if "時価総額" in display.columns:
    display["時価総額"] = display["時価総額"].map(format_jpy)

  for col in ["売上成長率", "純利益率", "ROE", "配当利回り", "52週高値比"]:
    if col in display.columns:
      display[col] = display[col].map(
          lambda x: f"{x * 100:.1f}%" if pd.notna(x) else "-"
      )

  for col in [
      "ベータ値",
      "実績PER",
      "PBR",
      "品質スコア",
      "割安スコア",
      "総合スコア",
  ]:
    if col in display.columns:
      display[col] = display[col].map(lambda x: f"{x:.2f}" if pd.notna(x) else "-")

  return display


# ============================================================
# スプレッドシート用エクスポート関数
# ============================================================
def create_spreadsheet_export_dataframe(dataframe):
  df = dataframe.copy()

  percent_columns = [
      "revenue_growth_1y",
      "revenue_cagr_3y",
      "operating_margin",
      "net_margin",
      "fcf_margin",
      "roe",
      "dividend_yield",
      "pct_from_52w_high",
  ]
  for col in percent_columns:
    if col in df.columns:
      df[col] = pd.to_numeric(df[col], errors="coerce") * 100

  cols_to_use = [c for c in JAPANESE_COLUMNS_MAP.keys() if c in df.columns]
  df = df[cols_to_use].rename(columns=JAPANESE_COLUMNS_MAP)
  return df.round(2)


def get_score_logic_dataframe():
  data = [
      {
          "カテゴリ": "クオリティ(60%配分)",
          "指標名": "売上高成長率(1年)",
          "評価基準": "高いほど高得点",
          "重み": "1.0",
          "概要": "前年比の売上成長の勢い",
      },
      {
          "カテゴリ": "クオリティ(60%配分)",
          "指標名": "売上高3年CAGR",
          "評価基準": "高いほど高得点",
          "重み": "1.0",
          "概要": "3年間の年平均成長率（中期的な成長力）",
      },
      {
          "カテゴリ": "クオリティ(60%配分)",
          "指標名": "純利益率",
          "評価基準": "高いほど高得点",
          "重み": "1.0",
          "概要": "売上に対する最終純利益の割合",
      },
      {
          "カテゴリ": "クオリティ(60%配分)",
          "指標名": "ROE(自己資本利益率)",
          "評価基準": "高いほど高得点",
          "重み": "1.0",
          "概要": "株主資本を使ってどれだけ利益を上げたか(日本株最重要)",
      },
      {
          "カテゴリ": "クオリティ(60%配分)",
          "指標名": "FCF利益率",
          "評価基準": "高いほど高得点",
          "重み": "1.0",
          "概要": "売上に対する自由に使える現金の割合",
      },
      {
          "カテゴリ": "クオリティ(60%配分)",
          "指標名": "ネット負債/EBITDA",
          "評価基準": "低いほど高得点",
          "重み": "0.5",
          "概要": "純有利子負債を本業利益で何年で返せるか(財務健全性)",
      },
      {
          "カテゴリ": "割安度(40%配分)",
          "指標名": "実績PER",
          "評価基準": "低いほど高得点(割安)",
          "重み": "1.0",
          "概要": "過去12ヶ月の純利益に対する株価倍率",
      },
      {
          "カテゴリ": "割安度(40%配分)",
          "指標名": "予想PER",
          "評価基準": "低いほど高得点(割安)",
          "重み": "1.0",
          "概要": "今期予想純利益に対する株価倍率",
      },
      {
          "カテゴリ": "割安度(40%配分)",
          "指標名": "PBR",
          "評価基準": "低いほど高得点(割安)",
          "重み": "0.5",
          "概要": "純資産に対する倍率(東証のPBR1倍割れ改善要請で注目)",
      },
      {
          "カテゴリ": "割安度(40%配分)",
          "指標名": "PSR",
          "評価基準": "低いほど高得点(割安)",
          "重み": "0.5",
          "概要": "売上高に対する株価倍率",
      },
      {
          "カテゴリ": "割安度(40%配分)",
          "指標名": "EV/EBITDA",
          "評価基準": "低いほど高得点(割安)",
          "重み": "1.0",
          "概要": "企業買収コストを本業利益で何年で回収できるか",
      },
      {
          "カテゴリ": "割安度(40%配分)",
          "指標名": "PEGレシオ",
          "評価基準": "低いほど高得点(割安)",
          "重み": "0.5",
          "概要": "PERを成長率で割った指標(成長に対して割安か)",
      },
      {
          "カテゴリ": "テクニカル指標",
          "指標名": "200日線上",
          "評価基準": "株価 >= 200日SMA",
          "重み": "-",
          "概要": "長期上昇トレンド判定（落ちてくるナイフの回避）",
      },
      {
          "カテゴリ": "テクニカル指標",
          "指標名": "52週高値乖離率",
          "評価基準": "0%に近いほど強い",
          "重み": "-",
          "概要": "新高値近辺のリーダー株を捉えるモメンタム指標",
      },
      {
          "カテゴリ": "リスク指標",
          "指標名": "ベータ値(Beta)",
          "評価基準": "<1.0=低変動, >1.2=高感応",
          "重み": "-",
          "概要": "日経平均に対する株価のボラティリティ感応度",
      },
  ]
  return pd.DataFrame(data)


def get_excel_download_with_autofilter(
    screened_jp, all_companies_jp, errors_df
):
  try:
    import openpyxl
    from openpyxl.utils import get_column_letter
  except ImportError:
    return None

  output = io.BytesIO()
  with pd.ExcelWriter(output, engine="openpyxl") as writer:
    all_companies_jp.to_excel(
        writer, sheet_name="日経225全銘柄(ソート済)", index=False
    )
    screened_jp.to_excel(writer, sheet_name="条件通過上位銘柄", index=False)
    score_logic_df = get_score_logic_dataframe()
    score_logic_df.to_excel(
        writer, sheet_name="総合スコアの算出基準", index=False
    )

    if not errors_df.empty:
      errors_df.to_excel(writer, sheet_name="取得エラー", index=False)

    workbook = writer.book
    for sheet_name in ["日経225全銘柄(ソート済)", "条件通過上位銘柄"]:
      if sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
          max_len = max(len(str(cell.value or "")) for cell in col)
          col_letter = get_column_letter(col[0].column)
          ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

  return output.getvalue()


# ============================================================
# Streamlit UI (メイン画面)
# ============================================================
def main():
  st.title("🇯🇵 日経225 高度スクリーニング")
  st.write(
      "日本を代表する日経平均株価（225銘柄）のファンダメンタルズ（稼ぐ力・割安性）とテクニカル（トレンド・モメンタム）を統合評価するスクリーニングツールです。"
  )

  with st.expander("ℹ️ 【解説】総合スコアおよび評価指標について"):
    st.markdown("""
        ### 📊 総合スコア (0〜100点満点)
        日経225銘柄の中で各指標を順位付け（パーセンタイル順位）し、**「企業の質（稼ぐ力・健全性）」**と**「株価の割安度」**を統合したスコアです。
        
        $$\\text{総合スコア} = (\\text{クオリティスコア} \\times 0.60) + (\\text{割安スコア} \\times 0.40)$$
        """)

    col_q, col_v = st.columns(2)
    with col_q:
      st.markdown("#### ① クオリティスコア (配分 60%)")
      st.caption("稼ぐ力・利益の質・財務健全性を評価（高いほど高得点）")
      st.markdown("""
            * **売上高成長率 (1年)** (重み: 1.0)
            * **売上高3年CAGR (年平均成長率)** (重み: 1.0)
            * **純利益率** (重み: 1.0)
            * **ROE (自己資本利益率)** (重み: 1.0) ※日本企業で最重要視
            * **FCF利益率 (フリーキャッシュフロー/売上高)** (重み: 1.0)
            * **ネット有利子負債 / EBITDA** (重み: 0.5) ※低いほど健全
            """)
    with col_v:
      st.markdown("#### ② 割安スコア (配分 40%)")
      st.caption("株価の割安度を評価（低いほど割安・高得点）")
      st.markdown("""
            * **実績PER** (重み: 1.0)
            * **予想PER** (重み: 1.0)
            * **PBR (株価純資産倍率)** (重み: 0.5) ※東証要請で注目の指標
            * **EV / EBITDA** (重み: 1.0)
            * **PSR (株価売上高倍率)** (重み: 0.5)
            * **PEGレシオ** (重み: 0.5)
            """)

    st.markdown("---")
    st.markdown("#### 🎯 日本株トレードにおける非財務フィルターの役割")
    st.markdown("""
        * **200日移動平均線（長期トレンド）**: 日本株でも「株価が200日移動平均線より上」にある銘柄に絞ることで、業績が良くても海外勢に売られている下落トレンド株（バリューの罠）を確実に排除できます。
        * **52週高値圏モメンタム**: 日経平均を牽引する最強の主導株を捉えます。
        * **配当利回り**: 日本株投資で人気の高い「高配当優良株」の発掘に活用できます。
        """)

  # session_state の初期化
  if "all_companies" not in st.session_state:
    st.session_state["all_companies"] = None
  if "errors" not in st.session_state:
    st.session_state["errors"] = []

  # 初回起動時にローカル保存データが存在すれば自動ロード
  if st.session_state["all_companies"] is None:
    saved_df = load_data_from_local()
    if saved_df is not None:
      st.session_state["all_companies"] = saved_df

  # ============================================================
  # サイドバー設定
  # ============================================================
  st.sidebar.header("📁 保存データの管理")
  saved_time = get_saved_data_time()
  if saved_time:
    st.sidebar.success(f"保存データあり\n(更新: {saved_time})")
    if st.sidebar.button("📂 保存データを再読み込み", use_container_width=True):
      st.session_state["all_companies"] = load_data_from_local()
      st.rerun()
  else:
    st.sidebar.info("保存データはまだありません。\n取得すると自動保存されます。")

  st.sidebar.markdown("---")
  st.sidebar.header("1. データ取得設定")
  max_tickers = st.sidebar.number_input(
      "新規取得する銘柄数",
      min_value=1,
      max_value=230,
      value=225,
      step=25,
      help="全225銘柄を取得する場合、約5〜8分で完了します。",
  )

  config = DEFAULT_CONFIG.copy()

  # データ新規取得ボタン
  if st.sidebar.button(
      "🔄 最新データを取得（上書き保存）",
      type="primary",
      use_container_width=True,
  ):
    constituents = get_nikkei225_constituents(config["nikkei225_url"])
    constituents = constituents.head(max_tickers)
    total = len(constituents)

    results, errors = [], []
    progress_bar = st.progress(0)
    status_text = st.empty()

    for number, (_, company) in enumerate(
        constituents.iterrows(), start=1
    ):
      symbol = company["display_symbol"]
      name = company["Security"]
      status_text.text(
          f"[{number:>3}/{total}] 財務・バリュエーション取得中：{name} ({symbol})"
          " ..."
      )

      try:
        result = analyze_company(company)
        results.append(result)
      except Exception as e:
        errors.append(
            {"symbol": symbol, "company_name": name, "error": str(e)}
        )

      time.sleep(config["request_sleep_seconds"])
      progress_bar.progress(number / total)

    if not results:
      st.error("有効なデータを取得できませんでした。")
      return

    status_text.info("スコア計算中...")
    all_companies = pd.DataFrame(results)
    all_companies = add_scores(all_companies, config)

    # 総合スコア順にソート
    sort_by = config.get("sort_by", "composite_score")
    sort_asc = config.get("sort_ascending", False)
    all_companies = all_companies.sort_values(
        sort_by, ascending=sort_asc, na_position="last"
    ).reset_index(drop=True)

    # data_japan/ フォルダへ自動保存
    save_data_to_local(all_companies)

    st.session_state["all_companies"] = all_companies
    st.session_state["errors"] = errors
    status_text.success(
        f"完了しました！（全{len(all_companies)}銘柄取得・data_japanフォルダに保存完了）"
    )

  st.sidebar.markdown("---")
  st.sidebar.header("2. 財務フィルター設定")
  use_financial_filter = st.sidebar.checkbox(
      "財務フィルターを適用する",
      value=False,
      help="ROE 8%以上、純利益率 8%以上、PER 35倍以下などの基準を適用します。",
  )

  st.sidebar.markdown("---")
  st.sidebar.header("3. 🎯 テクニカル・非財務フィルター")

  filter_sma200 = st.sidebar.checkbox(
      "長期上昇トレンドのみ (株価 > 200日線)",
      value=True,
      help="株価が200日移動平均線の上にある銘柄に限定し、下降トレンド株（バリューの罠）を排除します。",
  )

  filter_momentum = st.sidebar.checkbox(
      "高値圏モメンタム株 (52週高値から-20%以内)",
      value=False,
      help="直近52週高値から20%以内にある強い銘柄のみを抽出します。",
  )

  beta_option = st.sidebar.selectbox(
      "ベータ値（リスク感応度）",
      options=[
          "指定なし（全銘柄）",
          "低ボラ・守り重視 (Beta < 1.0)",
          "高ボラ・攻め重視 (Beta > 1.2)",
      ],
      index=0,
      help="日経平均に対する値動きの感応度を選択します。",
  )

  min_dividend = st.sidebar.slider(
      "最低配当利回り (%)",
      min_value=0.0,
      max_value=6.0,
      value=0.0,
      step=0.5,
      help="配当狙いの場合に設定します。0%は無配株も含みます。",
  )

  # セクターの動的取得
  available_sectors = []
  if st.session_state["all_companies"] is not None:
    available_sectors = sorted(
        st.session_state["all_companies"]["sector"].dropna().unique().tolist()
    )

  selected_sectors = st.sidebar.multiselect(
      "業種(セクター)絞り込み",
      options=available_sectors,
      default=available_sectors,
      help="選択した業種の銘柄のみが表示されます。",
  )

  st.sidebar.markdown("---")
  top_n = st.sidebar.number_input(
      "Webプレビューに表示する上位銘柄数",
      min_value=5,
      max_value=225,
      value=50,
      step=5,
      help="抽出された銘柄のうち、総合スコア上位何件を表示するか設定します。",
  )

  # ============================================================
  # 画面表示＆ダウンロード処理
  # ============================================================
  if st.session_state["all_companies"] is not None:
    all_companies = st.session_state["all_companies"]
    errors = st.session_state["errors"]

    filtered_df = all_companies.copy()

    # 1. 財務フィルター
    if use_financial_filter:
      filtered_df = apply_financial_filters(filtered_df, config)

    # 2. セクターフィルター
    if selected_sectors:
      filtered_df = filtered_df[filtered_df["sector"].isin(selected_sectors)]

    # 3. 200日線フィルター
    if filter_sma200 and "is_above_sma200" in filtered_df.columns:
      filtered_df = filtered_df[filtered_df["is_above_sma200"] == True]

    # 4. 52週高値モメンタムフィルター
    if filter_momentum and "pct_from_52w_high" in filtered_df.columns:
      filtered_df = filtered_df[filtered_df["pct_from_52w_high"] >= -0.20]

    # 5. ベータ値フィルター
    if beta_option == "低ボラ・守り重視 (Beta < 1.0)" and "beta" in filtered_df.columns:
      filtered_df = filtered_df[filtered_df["beta"] < 1.0]
    elif (
        beta_option == "高ボラ・攻め重視 (Beta > 1.2)"
        and "beta" in filtered_df.columns
    ):
      filtered_df = filtered_df[filtered_df["beta"] > 1.2]

    # 6. 配当利回りフィルター
    if min_dividend > 0 and "dividend_yield" in filtered_df.columns:
      filtered_df = filtered_df[
          filtered_df["dividend_yield"] >= (min_dividend / 100)
      ]

    # 上位N件の抽出
    sort_by = config.get("sort_by", "composite_score")
    sort_asc = config.get("sort_ascending", False)
    screened = (
        filtered_df.sort_values(
            sort_by, ascending=sort_asc, na_position="last"
        )
        .head(int(top_n))
        .reset_index(drop=True)
    )

    all_companies_jp = create_spreadsheet_export_dataframe(all_companies)
    screened_jp = create_spreadsheet_export_dataframe(screened)

    # ダウンロード＆保存先案内エリア
    st.markdown("---")
    st.subheader("📥 スプレッドシート用データダウンロード＆保存状況")

    saved_time_str = get_saved_data_time() or "未保存"
    st.info(
        f"💾 **ローカル保存状況**: `data_japan/` フォルダ内に保存済み（更新日時:"
        f" {saved_time_str}）\n\n※次回アプリ起動時は自動で前回のデータが読み込まれます。"
    )

    col1, col2 = st.columns(2)
    excel_data = get_excel_download_with_autofilter(
        screened_jp, all_companies_jp, pd.DataFrame(errors)
    )
    if excel_data is not None:
      col1.download_button(
          label="📗 ソート機能＆解説付きExcel (.xlsx) をダウンロード",
          data=excel_data,
          file_name="nikkei225_screening_autofilter.xlsx",
          mime=(
              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
          ),
          type="primary",
      )
    else:
      col1.warning("※Excel出力には `openpyxl` が必要です。")

    all_csv_data = all_companies_jp.to_csv(
        index=False, encoding="utf-8-sig"
    ).encode("utf-8-sig")
    col2.download_button(
        label=f"📄 日経225全{len(all_companies_jp)}銘柄データ (日本語CSV)",
        data=all_csv_data,
        file_name="nikkei225_all_companies_sorted.csv",
        mime="text/csv",
    )

    # Webプレビュー表示
    st.markdown("---")
    total_passed = len(filtered_df)

    st.subheader(
        f"🏆 条件合致 {total_passed} 銘柄（スコア上位 {len(screened)} 銘柄を表示中）"
    )

    if filtered_df.empty:
      st.warning(
          "⚠️ 設定した条件に合致する銘柄がありませんでした。サイドバーのフィルター条件を緩めてください。"
      )
    else:
      display_df = create_display_dataframe(screened)
      st.dataframe(display_df, use_container_width=True)

    if errors:
      with st.expander(f"⚠️ 取得エラー銘柄一覧 ({len(errors)}件)"):
        st.dataframe(pd.DataFrame(errors))
  else:
    st.warning(
        "👈 サイドバーの「🔄 最新データを取得（上書き保存）」ボタンを押してデータを取得してください。"
    )


if __name__ == "__main__":
  main()
