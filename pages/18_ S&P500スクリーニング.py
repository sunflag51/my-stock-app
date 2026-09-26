import io
import time
import warnings
import numpy as np
import pandas as pd
import streamlit as st
import yfinance as yf

warnings.filterwarnings("ignore")

# ============================================================
# 初期設定 (Streamlit画面の設定)
# ============================================================
st.set_page_config(page_title="S&P500 高度スクリーニング", layout="wide")

DEFAULT_CONFIG = {
    "sp500_url": "https://en.wikipedia.org/wiki/List_of_S%26P_500_companies",
    "top_n": 50,
    "request_sleep_seconds": 0.15,
    "missing_data_passes": False,
    "excluded_sectors": [],
    "filters": {
        "revenue_usd": {"min": 10_000_000_000, "max": None},
        "revenue_growth_1y": {"min": 0.10, "max": None},
        "revenue_cagr_3y": {"min": 0.08, "max": None},
        "net_margin": {"min": 0.15, "max": None},
        "roe": {"min": 0.15, "max": None},
        "fcf_margin": {"min": 0.08, "max": None},
        "net_debt_to_ebitda": {"min": None, "max": 2.50},
        "trailing_pe": {"min": 0, "max": 40},
        "forward_pe": {"min": 0, "max": 35},
        "price_to_book": {"min": None, "max": None},
        "price_to_sales": {"min": None, "max": 15},
        "ev_to_ebitda": {"min": 0, "max": 25},
        "peg_ratio": {"min": 0, "max": None},
    },
    "quality_score_metrics": {
        "revenue_growth_1y": {"direction": "high", "weight": 1.0},
        "revenue_cagr_3y": {"direction": "high", "weight": 1.0},
        "net_margin": {"direction": "high", "weight": 1.0},
        "roe": {"direction": "high", "weight": 1.0},
        "fcf_margin": {"direction": "high", "weight": 1.0},
        "net_debt_to_ebitda": {"direction": "low", "weight": 0.5},
    },
    "valuation_score_metrics": {
        "trailing_pe": {"direction": "low", "weight": 1.0},
        "forward_pe": {"direction": "low", "weight": 1.0},
        "price_to_book": {"direction": "low", "weight": 0.5},
        "price_to_sales": {"direction": "low", "weight": 0.5},
        "ev_to_ebitda": {"direction": "low", "weight": 1.0},
        "peg_ratio": {"direction": "low", "weight": 0.5},
    },
    "score_weights": {"quality": 0.60, "valuation": 0.40},
    "sort_by": "composite_score",
    "sort_ascending": False,
}

# ============================================================
# スプレッドシート用 日本語カラム対応表
# ============================================================
JAPANESE_COLUMNS_MAP = {
    "symbol": "ティッカー",
    "company_name": "企業名",
    "sector": "セクター",
    "market_cap_usd": "時価総額(USD)",
    "enterprise_value_usd": "企業価値EV(USD)",
    "revenue_usd": "売上高(USD)",
    "net_income_usd": "純利益(USD)",
    "ebitda_usd": "EBITDA(USD)",
    "free_cash_flow_usd": "FCF(USD)",
    "revenue_growth_1y": "売上高成長率(%)",
    "revenue_cagr_3y": "売上高3年CAGR(%)",
    "operating_margin": "営業利益率(%)",
    "net_margin": "純利益率(%)",
    "fcf_margin": "FCF利益率(%)",
    "roe": "ROE(%)",
    "total_debt_usd": "総負債(USD)",
    "cash_usd": "保有現金(USD)",
    "net_debt_usd": "ネット有利子負債(USD)",
    "net_debt_to_ebitda": "ネット負債/EBITDA(倍)",
    "trailing_pe": "実績PER(倍)",
    "forward_pe": "予想PER(倍)",
    "price_to_book": "PBR(倍)",
    "price_to_sales": "PSR(倍)",
    "ev_to_ebitda": "EV/EBITDA(倍)",
    "peg_ratio": "PEGレシオ",
    "dividend_yield": "配当利回り(%)",
    "quality_score": "クオリティスコア",
    "valuation_score": "割安スコア",
    "composite_score": "総合スコア",
}


# ============================================================
# データ取得系
# ============================================================
@st.cache_data(ttl=3600 * 24)
def get_sp500_constituents(url):
  storage_options = {
      "User-Agent": (
          "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML,"
          " like Gecko) Chrome/120.0.0.0 Safari/537.36"
      )
  }
  table = pd.read_html(url, storage_options=storage_options)[0]
  required_columns = ["Symbol", "Security", "GICS Sector"]
  table = table[required_columns].copy()
  table["yahoo_symbol"] = table["Symbol"].str.replace(".", "-", regex=False)
  table["display_symbol"] = table["Symbol"] + ".US"
  return table


def get_statement_series(statement, possible_names):
  if statement is None or statement.empty:
    return pd.Series(dtype="float64")
  for name in possible_names:
    if name in statement.index:
      values = pd.to_numeric(statement.loc[name], errors="coerce").dropna()
      return values.sort_index(ascending=False)
  return pd.Series(dtype="float64")


def first_valid_value(series):
  if series is None or len(series) == 0:
    return np.nan
  value = series.iloc[0]
  return np.nan if pd.isna(value) else float(value)


def safe_float(value):
  try:
    if value is None or pd.isna(value):
      return np.nan
    return float(value)
  except (TypeError, ValueError):
    return np.nan


def safe_divide(numerator, denominator):
  if pd.isna(numerator) or pd.isna(denominator) or denominator == 0:
    return np.nan
  return numerator / denominator


def get_first_info_value(info, names):
  for name in names:
    value = safe_float(info.get(name))
    if not pd.isna(value):
      return value
  return np.nan


def get_market_cap(ticker, info):
  try:
    value = safe_float(ticker.fast_info["market_cap"])
    if not pd.isna(value):
      return value
  except Exception:
    pass
  return get_first_info_value(info, ["marketCap"])


def analyze_company(company):
  yahoo_symbol = company["yahoo_symbol"]
  ticker = yf.Ticker(yahoo_symbol)

  income_statement = ticker.get_income_stmt(freq="yearly")
  balance_sheet = ticker.get_balance_sheet(freq="yearly")
  cash_flow = ticker.get_cash_flow(freq="yearly")

  try:
    info = ticker.get_info()
  except Exception:
    info = {}

  revenue = get_statement_series(
      income_statement, ["TotalRevenue", "OperatingRevenue"]
  )
  net_income = get_statement_series(income_statement, [
      "NetIncomeCommonStockholders",
      "NetIncome",
      "NetIncomeIncludingNoncontrollingInterests",
  ])
  ebitda = get_statement_series(income_statement, ["EBITDA", "NormalizedEBITDA"])
  operating_income = get_statement_series(
      income_statement, ["OperatingIncome"]
  )

  equity = get_statement_series(balance_sheet, [
      "StockholdersEquity",
      "CommonStockEquity",
      "TotalEquityGrossMinorityInterest",
  ])
  total_debt = get_statement_series(balance_sheet, ["TotalDebt"])
  cash = get_statement_series(balance_sheet, [
      "CashCashEquivalentsAndShortTermInvestments",
      "CashAndCashEquivalents",
      "CashFinancial",
  ])

  free_cash_flow = get_statement_series(cash_flow, ["FreeCashFlow"])
  operating_cash_flow = get_statement_series(
      cash_flow, ["OperatingCashFlow", "TotalCashFromOperatingActivities"]
  )
  capital_expenditure = get_statement_series(
      cash_flow, ["CapitalExpenditure", "CapitalExpenditures"]
  )

  latest_revenue = first_valid_value(revenue)
  latest_net_income = first_valid_value(net_income)
  latest_ebitda = first_valid_value(ebitda)
  latest_operating_income = first_valid_value(operating_income)
  latest_equity = first_valid_value(equity)
  latest_total_debt = first_valid_value(total_debt)
  latest_cash = first_valid_value(cash)
  latest_fcf = first_valid_value(free_cash_flow)
  latest_ocf = first_valid_value(operating_cash_flow)
  latest_capex = first_valid_value(capital_expenditure)

  if (
      pd.isna(latest_fcf)
      and not pd.isna(latest_ocf)
      and not pd.isna(latest_capex)
  ):
    latest_fcf = (
        latest_ocf + latest_capex
        if latest_capex < 0
        else latest_ocf - latest_capex
    )

  revenue_growth_1y = np.nan
  if len(revenue) >= 2:
    prev_rev = safe_float(revenue.iloc[1])
    if prev_rev > 0:
      revenue_growth_1y = latest_revenue / prev_rev - 1

  revenue_cagr_3y = np.nan
  if len(revenue) >= 4:
    rev_3y = safe_float(revenue.iloc[3])
    if latest_revenue > 0 and rev_3y > 0:
      revenue_cagr_3y = (latest_revenue / rev_3y) ** (1 / 3) - 1

  net_margin = safe_divide(latest_net_income, latest_revenue)
  operating_margin = safe_divide(latest_operating_income, latest_revenue)
  fcf_margin = safe_divide(latest_fcf, latest_revenue)

  roe = np.nan
  if len(equity) >= 2:
    prev_eq = safe_float(equity.iloc[1])
    if latest_equity > 0 and prev_eq > 0:
      roe = safe_divide(latest_net_income, (latest_equity + prev_eq) / 2)

  net_debt = np.nan
  net_debt_to_ebitda = np.nan
  if not pd.isna(latest_total_debt):
    cash_val = 0 if pd.isna(latest_cash) else latest_cash
    net_debt = latest_total_debt - cash_val
    if not pd.isna(latest_ebitda) and latest_ebitda > 0:
      net_debt_to_ebitda = net_debt / latest_ebitda

  market_cap = get_market_cap(ticker, info)

  return {
      "symbol": company["display_symbol"],
      "company_name": company["Security"],
      "sector": company["GICS Sector"],
      "market_cap_usd": market_cap,
      "enterprise_value_usd": get_first_info_value(info, ["enterpriseValue"]),
      "revenue_usd": latest_revenue,
      "net_income_usd": latest_net_income,
      "ebitda_usd": latest_ebitda,
      "free_cash_flow_usd": latest_fcf,
      "revenue_growth_1y": revenue_growth_1y,
      "revenue_cagr_3y": revenue_cagr_3y,
      "operating_margin": operating_margin,
      "net_margin": net_margin,
      "fcf_margin": fcf_margin,
      "roe": roe,
      "total_debt_usd": latest_total_debt,
      "cash_usd": latest_cash,
      "net_debt_usd": net_debt,
      "net_debt_to_ebitda": net_debt_to_ebitda,
      "trailing_pe": get_first_info_value(info, ["trailingPE"]),
      "forward_pe": get_first_info_value(info, ["forwardPE"]),
      "price_to_book": get_first_info_value(info, ["priceToBook"]),
      "price_to_sales": get_first_info_value(
          info, ["priceToSalesTrailing12Months"]
      ),
      "ev_to_ebitda": get_first_info_value(info, ["enterpriseToEbitda"]),
      "peg_ratio": get_first_info_value(info, ["trailingPegRatio", "pegRatio"]),
      "dividend_yield": get_first_info_value(info, ["dividendYield"]),
  }


# ============================================================
# スコア計算・フィルタリング
# ============================================================
def clean_valuation_values(dataframe):
  val_cols = [
      "trailing_pe",
      "forward_pe",
      "price_to_book",
      "price_to_sales",
      "ev_to_ebitda",
      "peg_ratio",
  ]
  for col in val_cols:
    if col in dataframe.columns:
      dataframe.loc[dataframe[col] <= 0, col] = np.nan
  return dataframe


def calculate_metric_score(series, direction):
  numeric = pd.to_numeric(series, errors="coerce")
  if direction == "high":
    return numeric.rank(pct=True, ascending=True) * 100
  elif direction == "low":
    return numeric.rank(pct=True, ascending=False) * 100
  return np.nan


def calculate_group_score(dataframe, metrics):
  weighted_scores, weights = [], []
  for metric, settings in metrics.items():
    if metric not in dataframe.columns:
      continue
    weight = float(settings.get("weight", 1.0))
    m_score = calculate_metric_score(
        dataframe[metric], settings.get("direction", "high")
    )
    weighted_scores.append(m_score * weight)
    weights.append(
        pd.Series(
            np.where(m_score.notna(), weight, 0), index=dataframe.index
        )
    )

  if not weighted_scores:
    return pd.Series(np.nan, index=dataframe.index)
  return sum(weighted_scores).divide(sum(weights).replace(0, np.nan))


def add_scores(dataframe, config):
  df = clean_valuation_values(dataframe.copy())
  df["quality_score"] = calculate_group_score(
      df, config["quality_score_metrics"]
  )
  df["valuation_score"] = calculate_group_score(
      df, config["valuation_score_metrics"]
  )

  q_weight = config["score_weights"]["quality"]
  v_weight = config["score_weights"]["valuation"]

  valid_q = np.where(df["quality_score"].notna(), q_weight, 0)
  valid_v = np.where(df["valuation_score"].notna(), v_weight, 0)

  df["composite_score"] = (
      (df["quality_score"] * q_weight).fillna(0)
      + (df["valuation_score"] * v_weight).fillna(0)
  ) / np.where((valid_q + valid_v) == 0, np.nan, (valid_q + valid_v))
  return df


def apply_filters(dataframe, config):
  filtered = dataframe.copy()
  if config.get("excluded_sectors"):
    filtered = filtered[~filtered["sector"].isin(config["excluded_sectors"])]

  for metric, limits in config["filters"].items():
    if metric not in filtered.columns:
      continue
    min_v, max_v = limits.get("min"), limits.get("max")
    values = pd.to_numeric(filtered[metric], errors="coerce")
    cond = pd.Series(True, index=filtered.index)

    if min_v is not None:
      cond &= values >= min_v
    if max_v is not None:
      cond &= values <= max_v

    if min_v is not None or max_v is not None:
      cond = (
          cond | values.isna()
          if config.get("missing_data_passes")
          else cond & values.notna()
      )
    filtered = filtered[cond]
  return filtered


def create_display_dataframe(dataframe):
  display_cols = [
      "symbol",
      "company_name",
      "sector",
      "market_cap_usd",
      "revenue_growth_1y",
      "revenue_cagr_3y",
      "net_margin",
      "roe",
      "fcf_margin",
      "net_debt_to_ebitda",
      "trailing_pe",
      "forward_pe",
      "price_to_book",
      "price_to_sales",
      "ev_to_ebitda",
      "peg_ratio",
      "quality_score",
      "valuation_score",
      "composite_score",
  ]
  display = dataframe[
      [c for c in display_cols if c in dataframe.columns]
  ].copy()

  rename_map = {
      "symbol": "銘柄",
      "company_name": "会社名",
      "sector": "セクター",
      "market_cap_usd": "時価総額",
      "revenue_growth_1y": "売上成長率",
      "revenue_cagr_3y": "売上3年CAGR",
      "net_margin": "純利益率",
      "roe": "ROE",
      "fcf_margin": "FCF利益率",
      "net_debt_to_ebitda": "ネット負債/EBITDA",
      "trailing_pe": "実績PER",
      "forward_pe": "予想PER",
      "price_to_book": "PBR",
      "price_to_sales": "PSR",
      "ev_to_ebitda": "EV/EBITDA",
      "peg_ratio": "PEG",
      "quality_score": "品質スコア",
      "valuation_score": "割安スコア",
      "composite_score": "総合スコア",
  }
  display = display.rename(columns=rename_map)

  if "時価総額" in display.columns:
    display["時価総額"] = display["時価総額"].map(
        lambda x: f"${x / 1_000_000_000:,.1f}B" if pd.notna(x) else "-"
    )

  for col in ["売上成長率", "売上3年CAGR", "純利益率", "ROE", "FCF利益率"]:
    if col in display.columns:
      display[col] = display[col].map(
          lambda x: f"{x * 100:.1f}%" if pd.notna(x) else "-"
      )

  for col in [
      "ネット負債/EBITDA",
      "実績PER",
      "予想PER",
      "PBR",
      "PSR",
      "EV/EBITDA",
      "PEG",
      "品質スコア",
      "割安スコア",
      "総合スコア",
  ]:
    if col in display.columns:
      display[col] = display[col].map(lambda x: f"{x:.2f}" if pd.notna(x) else "-")

  return display


# ============================================================
# スプレッドシート用エクスポート関数
# ============================================================
def create_spreadsheet_export_dataframe(dataframe):
  df = dataframe.copy()

  percent_columns = [
      "revenue_growth_1y",
      "revenue_cagr_3y",
      "operating_margin",
      "net_margin",
      "fcf_margin",
      "roe",
      "dividend_yield",
  ]
  for col in percent_columns:
    if col in df.columns:
      df[col] = df[col] * 100

  cols_to_use = [c for c in JAPANESE_COLUMNS_MAP.keys() if c in df.columns]
  df = df[cols_to_use].rename(columns=JAPANESE_COLUMNS_MAP)
  return df.round(2)


def get_score_logic_dataframe():
  data = [
      {
          "カテゴリ": "クオリティ(60%配分)",
          "指標名": "売上高成長率(1年)",
          "評価基準": "高いほど高得点",
          "重み": "1.0",
          "概要": "前年比の売上成長の勢い",
      },
      {
          "カテゴリ": "クオリティ(60%配分)",
          "指標名": "売上高3年CAGR",
          "評価基準": "高いほど高得点",
          "重み": "1.0",
          "概要": "3年間の年平均成長率（中期的な成長力）",
      },
      {
          "カテゴリ": "クオリティ(60%配分)",
          "指標名": "純利益率",
          "評価基準": "高いほど高得点",
          "重み": "1.0",
          "概要": "売上に対する最終純利益の割合",
      },
      {
          "カテゴリ": "クオリティ(60%配分)",
          "指標名": "ROE(自己資本利益率)",
          "評価基準": "高いほど高得点",
          "重み": "1.0",
          "概要": "株主資本を使ってどれだけ利益を上げたか",
      },
      {
          "カテゴリ": "クオリティ(60%配分)",
          "指標名": "FCF利益率",
          "評価基準": "高いほど高得点",
          "重み": "1.0",
          "概要": "売上に対する自由に使える現金の割合",
      },
      {
          "カテゴリ": "クオリティ(60%配分)",
          "指標名": "ネット負債/EBITDA",
          "評価基準": "低いほど高得点",
          "重み": "0.5",
          "概要": "純有利子負債を本業利益で何年で返せるか(健全性)",
      },
      {
          "カテゴリ": "割安度(40%配分)",
          "指標名": "実績PER",
          "評価基準": "低いほど高得点(割安)",
          "重み": "1.0",
          "概要": "過去12ヶ月の純利益に対する株価倍率",
      },
      {
          "カテゴリ": "割安度(40%配分)",
          "指標名": "予想PER",
          "評価基準": "低いほど高得点(割安)",
          "重み": "1.0",
          "概要": "将来予想利益に対する株価倍率",
      },
      {
          "カテゴリ": "割安度(40%配分)",
          "指標名": "EV/EBITDA",
          "評価基準": "低いほど高得点(割安)",
          "重み": "1.0",
          "概要": "企業買収コストを本業利益で何年で回収できるか",
      },
      {
          "カテゴリ": "割安度(40%配分)",
          "指標名": "PBR",
          "評価基準": "低いほど高得点(割安)",
          "重み": "0.5",
          "概要": "純資産に対する株価倍率",
      },
      {
          "カテゴリ": "割安度(40%配分)",
          "指標名": "PSR",
          "評価基準": "低いほど高得点(割安)",
          "重み": "0.5",
          "概要": "売上高に対する株価倍率",
      },
      {
          "カテゴリ": "割安度(40%配分)",
          "指標名": "PEGレシオ",
          "評価基準": "低いほど高得点(割安)",
          "重み": "0.5",
          "概要": "PERを成長率で割った指標(成長に対して割安か)",
      },
  ]
  return pd.DataFrame(data)


def get_excel_download_with_autofilter(
    screened_jp, all_companies_jp, errors_df
):
  try:
    import openpyxl
    from openpyxl.utils import get_column_letter
  except ImportError:
    return None

  output = io.BytesIO()
  with pd.ExcelWriter(output, engine="openpyxl") as writer:
    all_companies_jp.to_excel(
        writer, sheet_name="S&P500全銘柄(ソート済)", index=False
    )
    screened_jp.to_excel(writer, sheet_name="条件通過上位銘柄", index=False)
    score_logic_df = get_score_logic_dataframe()
    score_logic_df.to_excel(
        writer, sheet_name="総合スコアの算出基準", index=False
    )

    if not errors_df.empty:
      errors_df.to_excel(writer, sheet_name="取得エラー", index=False)

    workbook = writer.book
    for sheet_name in ["S&P500全銘柄(ソート済)", "条件通過上位銘柄"]:
      if sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
          max_len = max(len(str(cell.value or "")) for cell in col)
          col_letter = get_column_letter(col[0].column)
          ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

  return output.getvalue()


# ============================================================
# Streamlit UI (メイン画面)
# ============================================================
def main():
  st.title("🦅 S&P500 高度スクリーニング")
  st.write(
      "クオリティ（成長性・収益性）とバリュエーション（割安性）を統合したスコア評価とフィルタリングを行います。"
  )

  with st.expander("ℹ️ 【解説】総合スコアの算出ロジック・計算基準について"):
    st.markdown("""
        ### 📊 総合スコアの仕組み (0〜100点満点)
        全銘柄の中で各数値を順位付け（パーセンタイル順位）し、**「企業の質（稼ぐ力・健全性）」**と**「株価の割安度」**を統合したスコアです。
        
        $$\\text{総合スコア} = (\\text{クオリティスコア} \\times 0.60) + (\\text{割安スコア} \\times 0.40)$$
        """)

    col_q, col_v = st.columns(2)
    with col_q:
      st.markdown("#### ① クオリティスコア (配分 60%)")
      st.caption("稼ぐ力・利益の質・財務健全性を評価（高いほど高得点）")
      st.markdown("""
            * **売上高成長率 (1年)** (重み: 1.0)
            * **売上高3年CAGR (年平均成長率)** (重み: 1.0)
            * **純利益率** (重み: 1.0)
            * **ROE (自己資本利益率)** (重み: 1.0)
            * **FCF利益率 (フリーキャッシュフロー/売上高)** (重み: 1.0)
            * **ネット有利子負債 / EBITDA** (重み: 0.5) ※低いほど高得点
            """)
    with col_v:
      st.markdown("#### ② 割安スコア (配分 40%)")
      st.caption("株価の割安度を評価（低いほど高得点・割安）")
      st.markdown("""
            * **実績PER** (重み: 1.0)
            * **予想PER** (重み: 1.0)
            * **EV / EBITDA** (重み: 1.0)
            * **PBR (株価純資産倍率)** (重み: 0.5)
            * **PSR (株価売上高倍率)** (重み: 0.5)
            * **PEGレシオ** (重み: 0.5)
            """)
    st.info(
        "💡"
        " 赤字企業のマイナスPERなどの異常値は自動除外され、全銘柄中での相対順位で点数化されます。"
    )

  # session_state の初期化
  if "all_companies" not in st.session_state:
    st.session_state["all_companies"] = None
  if "errors" not in st.session_state:
    st.session_state["errors"] = []

  # サイドバー設定
  st.sidebar.header("実行設定")
  max_tickers = st.sidebar.number_input(
      "取得・分析する銘柄数",
      min_value=1,
      max_value=510,
      value=503,
      step=50,
      help="500銘柄すべて取得する場合は約15〜20分かかります。テスト時は少なめに設定してください。",
  )

  st.sidebar.markdown("---")
  st.sidebar.header("表示・絞り込み設定")

  # 【追加】フィルター適用のON/OFF
  use_filter = st.sidebar.checkbox(
      "財務フィルターを適用する",
      value=True,
      help=(
          "チェックを外すと、厳しい財務条件による足切りを行わず、S&P500全銘柄から純粋に総合スコアが高い順に表示します。"
      ),
  )

  # 表示件数の設定（即時反映）
  top_n = st.sidebar.number_input(
      "Webプレビューに表示する上位銘柄数",
      min_value=5,
      max_value=200,
      value=50,
      step=5,
      help="上位何件を表示するか設定します（データ取得後も即座に変更が反映されます）。",
  )

  config = DEFAULT_CONFIG.copy()

  # データ取得ボタン
  if st.button("データ取得＆スクリーニングを開始", type="primary"):
    constituents = get_sp500_constituents(config["sp500_url"])
    constituents = constituents.head(max_tickers)
    total = len(constituents)

    results, errors = [], []
    progress_bar = st.progress(0)
    status_text = st.empty()

    for number, (_, company) in enumerate(
        constituents.iterrows(), start=1
    ):
      symbol = company["display_symbol"]
      status_text.text(
          f"[{number:>3}/{total}] 財務・バリュエーション取得中：{symbol} ..."
      )

      try:
        result = analyze_company(company)
        results.append(result)
      except Exception as e:
        errors.append({
            "symbol": symbol,
            "company_name": company["Security"],
            "error": str(e),
        })

      time.sleep(config["request_sleep_seconds"])
      progress_bar.progress(number / total)

    if not results:
      st.error("有効なデータを取得できませんでした。")
      return

    status_text.info("スコア計算中...")
    all_companies = pd.DataFrame(results)
    all_companies = add_scores(all_companies, config)

    # 総合スコア順にソート
    sort_by = config.get("sort_by", "composite_score")
    sort_asc = config.get("sort_ascending", False)
    all_companies = all_companies.sort_values(
        sort_by, ascending=sort_asc, na_position="last"
    ).reset_index(drop=True)

    # session_state に全件データを保存
    st.session_state["all_companies"] = all_companies
    st.session_state["errors"] = errors

    status_text.success(f"完了しました！（全{len(all_companies)}銘柄取得）")

  # ============================================================
  # 画面表示＆ダウンロード（サイドバーの変更をリアルタイム反映）
  # ============================================================
  if st.session_state["all_companies"] is not None:
    all_companies = st.session_state["all_companies"]
    errors = st.session_state["errors"]

    # フィルターの適用判定（即時反映）
    if use_filter:
      passed_candidates = apply_filters(all_companies, config)
      filter_desc = "財務フィルター適用"
    else:
      passed_candidates = all_companies
      filter_desc = "全銘柄・スコア順"

    # 上位N件の抽出
    sort_by = config.get("sort_by", "composite_score")
    sort_asc = config.get("sort_ascending", False)
    passed_sorted = passed_candidates.sort_values(
        sort_by, ascending=sort_asc, na_position="last"
    ).reset_index(drop=True)

    # 指定された件数で切り出し
    screened = passed_sorted.head(int(top_n)).reset_index(drop=True)

    # スプレッドシート用データ作成
    all_companies_jp = create_spreadsheet_export_dataframe(all_companies)
    screened_jp = create_spreadsheet_export_dataframe(screened)

    # ダウンロードエリア
    st.markdown("---")
    st.subheader("📥 スプレッドシート用データダウンロード")

    col1, col2 = st.columns(2)
    excel_data = get_excel_download_with_autofilter(
        screened_jp, all_companies_jp, pd.DataFrame(errors)
    )
    if excel_data is not None:
      col1.download_button(
          label="📗 ソート機能＆解説付きExcel (.xlsx) をダウンロード",
          data=excel_data,
          file_name="sp500_screening_autofilter.xlsx",
          mime=(
              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
          ),
          type="primary",
      )
    else:
      col1.warning(
          "※Excel出力には `openpyxl` が必要です。右側のCSVをご利用ください。"
      )

    all_csv_data = all_companies_jp.to_csv(
        index=False, encoding="utf-8-sig"
    ).encode("utf-8-sig")
    col2.download_button(
        label=f"📄 S&P500全{len(all_companies_jp)}銘柄データ (日本語CSV)",
        data=all_csv_data,
        file_name="sp500_all_companies_sorted.csv",
        mime="text/csv",
    )

    # Webプレビュー表示エリア
    st.markdown("---")
    total_passed = len(passed_candidates)

    if use_filter and total_passed < top_n:
      st.subheader(
          f"🏆 条件通過銘柄（{total_passed}銘柄 / 設定上限: {top_n}件）"
      )
      st.info(
          f"💡 **条件を通過した銘柄が合計 {total_passed} 銘柄のため、全 {total_passed} 銘柄を表示しています。**\n"
          f"50銘柄以上表示したい場合は、サイドバーの **「財務フィルターを適用する」のチェックを外す** か、フィルター条件を緩和してください。"
      )
    else:
      st.subheader(
          f"🏆 スコア上位 {len(screened)} 銘柄（{filter_desc}・Webプレビュー）"
      )

    if screened.empty:
      st.warning("設定したフィルタ条件に合致する銘柄がありませんでした。")
    else:
      display_df = create_display_dataframe(screened)
      st.dataframe(display_df, use_container_width=True)

    if errors:
      with st.expander(f"⚠️ 取得エラー銘柄一覧 ({len(errors)}件)"):
        st.dataframe(pd.DataFrame(errors))


if __name__ == "__main__":
  main()
